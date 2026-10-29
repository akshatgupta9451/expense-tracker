"""Transaction reports as CSV, Excel (.xlsx) and PDF.

Exports are built in memory and returned directly. A row cap keeps that safe; very large or
scheduled reports would move to a background job (see the architecture notes).
"""

import csv
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from app.exceptions import AppError
from app.models import Transaction, TransactionType, User
from app.schemas.export import ExportFormat, ExportQuery
from app.services.transaction_service import count_transactions, filtered_transactions
from app.utils.money import ZERO

MAX_EXPORT_ROWS = 10_000

MEDIA_TYPES = {
    ExportFormat.CSV: "text/csv; charset=utf-8",
    ExportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ExportFormat.PDF: "application/pdf",
}

COLUMNS = [
    "Date",
    "Title",
    "Type",
    "Category",
    "Amount",
    "Currency",
    "Payment method",
    "Description",
]


@dataclass
class ExportFile:
    content: bytes
    filename: str
    media_type: str


@dataclass
class ReportData:
    user: User
    rows: list[Transaction]
    query: ExportQuery
    generated_at: datetime

    @property
    def income(self) -> Decimal:
        return sum((t.amount for t in self.rows if t.type == TransactionType.INCOME), ZERO)

    @property
    def expense(self) -> Decimal:
        return sum((t.amount for t in self.rows if t.type == TransactionType.EXPENSE), ZERO)

    def expense_by_category(self) -> list[tuple[str, Decimal]]:
        totals: dict[str, Decimal] = defaultdict(lambda: ZERO)
        for t in self.rows:
            if t.type == TransactionType.EXPENSE:
                totals[t.category.name] += t.amount
        return sorted(totals.items(), key=lambda kv: kv[1], reverse=True)

    def period_label(self) -> str:
        q = self.query
        if q.start_date and q.end_date:
            return f"{q.start_date:%b %d, %Y} – {q.end_date:%b %d, %Y}"
        if q.start_date:
            return f"From {q.start_date:%b %d, %Y}"
        if q.end_date:
            return f"Up to {q.end_date:%b %d, %Y}"
        if self.rows:
            dates = [t.transaction_date for t in self.rows]
            return f"{min(dates):%b %d, %Y} – {max(dates):%b %d, %Y}"
        return "All time"


def safe_cell(value: str | None) -> str:
    """Neutralise spreadsheet formula injection (CSV/Excel run cells starting with =, +, -, @)."""
    if not value:
        return value or ""
    return "'" + value if value[0] in "=+-@\t\r" else value


def _label(enum_value) -> str:
    return enum_value.value.replace("_", " ").capitalize() if enum_value else ""


def _filename(data: ReportData, fmt: ExportFormat) -> str:
    q = data.query
    if q.start_date or q.end_date:
        span = f"{q.start_date or 'start'}_to_{q.end_date or 'today'}"
    else:
        span = data.generated_at.strftime("%Y-%m-%d")
    return f"transactions_{span}.{fmt.value}"


# ---------- CSV ----------


def _csv(data: ReportData) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(COLUMNS)
    for t in data.rows:
        writer.writerow(
            [
                t.transaction_date.isoformat(),
                safe_cell(t.title),
                _label(t.type),
                safe_cell(t.category.name),
                f"{t.amount:.2f}" if t.type == TransactionType.INCOME else f"-{t.amount:.2f}",
                t.currency,
                _label(t.payment_method),
                safe_cell(t.description),
            ]
        )
    # UTF-8 BOM so Excel opens accented characters correctly.
    return "﻿".encode() + buf.getvalue().encode("utf-8")


# ---------- Excel ----------

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _money_format(currency: str) -> str:
    return f'#,##0.00 "{currency}";[Red]-#,##0.00 "{currency}"'


def _xlsx(data: ReportData) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for t in data.rows:
        signed = t.amount if t.type == TransactionType.INCOME else -t.amount
        ws.append(
            [
                t.transaction_date,
                safe_cell(t.title),
                _label(t.type),
                safe_cell(t.category.name),
                float(signed),
                t.currency,
                _label(t.payment_method),
                safe_cell(t.description),
            ]
        )
        row = ws.max_row
        ws.cell(row, 1).number_format = "yyyy-mm-dd"
        ws.cell(row, 5).number_format = _money_format(t.currency)

    widths = [12, 32, 10, 20, 16, 10, 16, 48]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"

    # Summary sheet
    s = wb.create_sheet("Summary")
    currency = data.user.currency
    s.append(["Transaction report"])
    s["A1"].font = Font(bold=True, size=14)
    for label, value in [
        ("Account", data.user.email),
        ("Period", data.period_label()),
        ("Generated", data.generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Transactions", len(data.rows)),
    ]:
        s.append([label, value])
    s.append([])
    for label, value in [
        ("Total income", data.income),
        ("Total expense", data.expense),
        ("Net", data.income - data.expense),
    ]:
        s.append([label, float(value)])
        s.cell(s.max_row, 1).font = Font(bold=True)
        s.cell(s.max_row, 2).number_format = _money_format(currency)
    s.append([])
    s.append(["Expenses by category", "Amount", "Share"])
    for cell in s[s.max_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    expense = data.expense
    for name, total in data.expense_by_category():
        s.append([safe_cell(name), float(total), float(total / expense) if expense else 0])
        s.cell(s.max_row, 2).number_format = _money_format(currency)
        s.cell(s.max_row, 3).number_format = "0.0%"
    s.column_dimensions["A"].width = 26
    s.column_dimensions["B"].width = 30
    s.column_dimensions["C"].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

INDIGO = colors.HexColor("#4F46E5")
GRAY = colors.HexColor("#6B7280")
LIGHT = colors.HexColor("#F3F4F6")
GREEN = colors.HexColor("#059669")
RED = colors.HexColor("#E11D48")

# Built-in PDF fonts only cover Latin-1/WinAnsi, so other currency symbols fall back to the code.
PDF_SYMBOLS = {"USD": "$", "EUR": "€", "GBP": "£", "JPY": "¥"}


def _pdf_money(value: Decimal, currency: str) -> str:
    symbol = PDF_SYMBOLS.get(currency)
    sign = "-" if value < 0 else ""
    amount = f"{abs(value):,.2f}"
    return f"{sign}{symbol}{amount}" if symbol else f"{sign}{amount} {currency}"


def _pdf(data: ReportData) -> bytes:
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title="Transaction report",
        author=data.user.name or data.user.email,
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle(
        "h1", parent=styles["Heading1"], fontSize=18, textColor=colors.black, spaceAfter=2
    )
    muted = ParagraphStyle("muted", parent=styles["Normal"], fontSize=9, textColor=GRAY)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=6)
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8.5, leading=10.5)
    cell_right = ParagraphStyle("cellr", parent=cell, alignment=TA_RIGHT)

    currency = data.user.currency
    money = lambda v: _pdf_money(v, currency)  # noqa: E731
    net = data.income - data.expense

    story = [
        Paragraph("Transaction report", h1),
        Paragraph(
            f"{_esc(data.user.name or data.user.email)} · {data.period_label()} · "
            f"generated {data.generated_at:%b %d, %Y %H:%M} UTC",
            muted,
        ),
        Spacer(1, 8),
    ]

    # Summary tiles
    tiles = [
        ("Income", money(data.income), GREEN),
        ("Expenses", money(data.expense), RED),
        ("Net", money(net), GREEN if net >= 0 else RED),
        ("Transactions", str(len(data.rows)), colors.black),
    ]
    label_style = ParagraphStyle("tl", parent=muted, fontSize=8)
    tile_cells = [
        [
            Paragraph(label.upper(), label_style),
            Paragraph(f'<font size="14" color="{c.hexval()}"><b>{v}</b></font>', cell),
        ]
        for label, v, c in tiles
    ]
    summary = Table([[t for t in tile_cells]], colWidths=[(page[0] - 28 * mm) / 4] * 4)
    summary.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), LIGHT),
                ("BOX", (0, 0), (-1, -1), 0, colors.white),
                ("LINEAFTER", (0, 0), (-2, -1), 6, colors.white),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    story.append(summary)

    # Category breakdown
    categories = data.expense_by_category()
    if categories:
        story.append(Paragraph("Expenses by category", h2))
        rows = [["Category", "Amount", "Share"]] + [
            [_esc(name), money(total), f"{(total / data.expense * 100):.1f}%"]
            for name, total in categories
        ]
        t = Table(rows, colWidths=[90 * mm, 40 * mm, 25 * mm], hAlign="LEFT")
        t.setStyle(_table_style(right_cols=(1, 2)))
        story.append(t)

    # Transactions
    story.append(Paragraph("Transactions", h2))
    if data.rows:
        rows = [["Date", "Title", "Category", "Type", "Payment", "Amount"]]
        for tx in data.rows:
            income = tx.type == TransactionType.INCOME
            amount = _pdf_money(tx.amount if income else -tx.amount, tx.currency)
            amount_color = (GREEN if income else colors.black).hexval()
            title = _esc(tx.title)
            if tx.description:
                title += (
                    f'<br/><font color="#6B7280" size="7.5">{_esc(tx.description[:160])}</font>'
                )
            rows.append(
                [
                    tx.transaction_date.strftime("%Y-%m-%d"),
                    Paragraph(title, cell),
                    Paragraph(_esc(tx.category.name), cell),
                    _label(tx.type),
                    _label(tx.payment_method),
                    Paragraph(
                        f'<font color="{amount_color}">{amount}</font>',
                        cell_right,
                    ),
                ]
            )
        table = Table(
            rows,
            colWidths=[24 * mm, 104 * mm, 45 * mm, 22 * mm, 30 * mm, 34 * mm],
            repeatRows=1,
        )
        table.setStyle(_table_style(right_cols=(5,)))
        story.append(table)
    else:
        story.append(Paragraph("No transactions match these filters.", muted))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(GRAY)
        canvas.drawString(14 * mm, 9 * mm, "Expense Tracker")
        canvas.drawRightString(page[0] - 14 * mm, 9 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _table_style(right_cols=()) -> TableStyle:
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), INDIGO),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for col in right_cols:
        style.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    return TableStyle(style)


def _esc(text: str) -> str:
    """Escape user text for ReportLab's mini-markup."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------- entry point ----------

RENDERERS = {ExportFormat.CSV: _csv, ExportFormat.XLSX: _xlsx, ExportFormat.PDF: _pdf}


async def export_transactions(db: AsyncSession, user: User, query: ExportQuery) -> ExportFile:
    stmt = filtered_transactions(user.id, query)
    total = await count_transactions(db, stmt)
    if total > MAX_EXPORT_ROWS:
        raise AppError(
            f"That's {total:,} transactions. Exports are limited to {MAX_EXPORT_ROWS:,} at a time; "
            "narrow the date range or filters.",
            details={"total": total, "limit": MAX_EXPORT_ROWS},
        )
    rows = list((await db.scalars(stmt)).unique().all())
    data = ReportData(user=user, rows=rows, query=query, generated_at=datetime.now(UTC))

    # File generation is CPU-bound; keep it off the event loop.
    content = await run_in_threadpool(RENDERERS[query.format], data)
    return ExportFile(
        content=content,
        filename=_filename(data, query.format),
        media_type=MEDIA_TYPES[query.format],
    )
