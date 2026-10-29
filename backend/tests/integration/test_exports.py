import csv
import io

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services import export_service
from tests.conftest import register_and_login


@pytest.fixture
async def cats(auth_client):
    return {c["name"]: c["id"] for c in (await auth_client.get("/categories")).json()}


@pytest.fixture
async def data(auth_client, cats):
    async def add(category, title, amount, day, **extra):
        r = await auth_client.post(
            "/transactions",
            json={
                "category_id": cats[category],
                "title": title,
                "amount": amount,
                "transaction_date": day,
                **extra,
            },
        )
        assert r.status_code == 201, r.text

    await add("Salary", "October salary", "5000", "2026-10-01")
    await add("Housing & Rent", "Rent", "1500", "2026-10-02", payment_method="bank")
    await add("Groceries", "Café & crème", "120.50", "2026-10-05", description="Weekly shop")
    await add("Groceries", "Old groceries", "80", "2026-09-20")


async def _export(client, fmt, **params):
    r = await client.get("/exports/transactions", params={"format": fmt, **params})
    assert r.status_code == 200, r.text
    return r


def _csv_rows(r):
    text = r.content.decode("utf-8")
    assert text.startswith("﻿")  # BOM for Excel
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


# ---------- CSV ----------


async def test_csv_contents_and_headers(auth_client, data):
    r = await _export(auth_client, "csv", start_date="2026-10-01", end_date="2026-10-31")
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"] == (
        'attachment; filename="transactions_2026-10-01_to_2026-10-31.csv"'
    )
    assert r.headers["cache-control"] == "no-store"

    rows = _csv_rows(r)
    assert rows[0] == export_service.COLUMNS
    assert rows[1:] == [
        ["2026-10-05", "Café & crème", "Expense", "Groceries", "-120.50", "USD", "", "Weekly shop"],
        ["2026-10-02", "Rent", "Expense", "Housing & Rent", "-1500.00", "USD", "Bank", ""],
        ["2026-10-01", "October salary", "Income", "Salary", "5000.00", "USD", "", ""],
    ]


async def test_export_respects_filters_and_sort(auth_client, data):
    rows = _csv_rows(await _export(auth_client, "csv", type="expense", sort="amount"))
    assert [r[1] for r in rows[1:]] == ["Old groceries", "Café & crème", "Rent"]
    rows = _csv_rows(await _export(auth_client, "csv", q="rent"))
    assert [r[1] for r in rows[1:]] == ["Rent"]


async def test_formula_injection_is_neutralised(auth_client, cats):
    for title in ['=HYPERLINK("http://evil","click")', "+SUM(A1)", "-2+3", "@cmd"]:
        await auth_client.post(
            "/transactions",
            json={
                "category_id": cats["Groceries"],
                "title": title,
                "amount": "1",
                "transaction_date": "2026-10-01",
            },
        )
    csv_titles = [r[1] for r in _csv_rows(await _export(auth_client, "csv"))[1:]]
    assert all(t.startswith("'") for t in csv_titles), csv_titles

    wb = load_workbook(io.BytesIO((await _export(auth_client, "xlsx")).content))
    cells = [c.value for c in wb["Transactions"]["B"][1:]]
    assert all(isinstance(v, str) and v.startswith("'") for v in cells), cells
    assert all(c.data_type == "s" for c in wb["Transactions"]["B"][1:])  # strings, not formulas


# ---------- Excel ----------


async def test_xlsx_has_typed_cells_and_summary(auth_client, data):
    r = await _export(auth_client, "xlsx", start_date="2026-10-01", end_date="2026-10-31")
    assert r.headers["content-type"] == export_service.MEDIA_TYPES["xlsx"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Transactions", "Summary"]

    ws = wb["Transactions"]
    assert [c.value for c in ws[1]] == export_service.COLUMNS
    first = [c.value for c in ws[2]]
    assert first[0].date().isoformat() == "2026-10-05"  # a real date, not text
    assert first[4] == -120.5  # a real number
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:H4"

    summary = {
        row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row and row[0]
    }
    assert summary["Account"] == "alice@example.com"
    assert summary["Transactions"] == 3
    assert (summary["Total income"], summary["Total expense"], summary["Net"]) == (
        5000,
        1620.5,
        3379.5,
    )
    assert summary["Housing & Rent"] == 1500


# ---------- PDF ----------


async def test_pdf_report(auth_client, data):
    r = await _export(auth_client, "pdf", start_date="2026-10-01", end_date="2026-10-31")
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-")
    reader = PdfReader(io.BytesIO(r.content))
    assert reader.metadata.title == "Transaction report"
    text = " ".join(page.extract_text() for page in reader.pages)
    for expected in [
        "Transaction report",
        "Oct 01, 2026 – Oct 31, 2026",
        "$5,000.00",  # income tile
        "$1,620.50",  # expense tile
        "$3,379.50",  # net
        "Expenses by category",
        "Café & crème",
        "Weekly shop",
        "Page 1",
    ]:
        assert expected in text, expected
    assert "Old groceries" not in text  # outside the date range


async def test_pdf_paginates_and_escapes_markup(auth_client, cats, monkeypatch):
    for i in range(60):
        await auth_client.post(
            "/transactions",
            json={
                "category_id": cats["Groceries"],
                "title": f"<b>Item {i}</b> & co",
                "amount": "1",
                "transaction_date": "2026-10-01",
            },
        )
    reader = PdfReader(io.BytesIO((await _export(auth_client, "pdf")).content))
    assert len(reader.pages) > 1
    text = " ".join(p.extract_text() for p in reader.pages)
    assert "<b>Item 0</b> & co" in text  # shown literally, not interpreted
    assert f"Page {len(reader.pages)}" in text


async def test_empty_export_still_produces_files(auth_client):
    assert len(_csv_rows(await _export(auth_client, "csv"))) == 1
    text = (
        PdfReader(io.BytesIO((await _export(auth_client, "pdf")).content)).pages[0].extract_text()
    )
    assert "No transactions match these filters." in text


async def test_non_latin_currency_falls_back_to_code_in_pdf(auth_client, cats):
    await auth_client.patch("/users/me", json={"currency": "INR"})
    await auth_client.post(
        "/transactions",
        json={
            "category_id": cats["Salary"],
            "title": "Pay",
            "amount": "1000",
            "transaction_date": "2026-10-01",
        },
    )
    text = (
        PdfReader(io.BytesIO((await _export(auth_client, "pdf")).content)).pages[0].extract_text()
    )
    assert "1,000.00 INR" in text


# ---------- limits & security ----------


async def test_row_limit(auth_client, data, monkeypatch):
    monkeypatch.setattr(export_service, "MAX_EXPORT_ROWS", 2)
    r = await auth_client.get("/exports/transactions", params={"format": "csv"})
    assert r.status_code == 400
    assert r.json()["error"]["details"] == {"total": 4, "limit": 2}


async def test_bad_format_and_auth(client, auth_client):
    r = await auth_client.get("/exports/transactions", params={"format": "docx"})
    assert r.status_code == 422
    client.headers.pop("Authorization")
    assert (await client.get("/exports/transactions")).status_code == 401


async def test_only_own_transactions_are_exported(client, seeded):
    await register_and_login(client, email="owner@example.com")
    salary = next(
        c["id"] for c in (await client.get("/categories")).json() if c["name"] == "Salary"
    )
    await client.post(
        "/transactions",
        json={
            "category_id": salary,
            "title": "Secret bonus",
            "amount": "1",
            "transaction_date": "2026-10-01",
        },
    )
    await register_and_login(client, email="other@example.com")
    # Excel and PDF are compressed, so inspect the parsed contents, not the raw bytes.
    assert len(_csv_rows(await _export(client, "csv"))) == 1
    wb = load_workbook(io.BytesIO((await _export(client, "xlsx")).content))
    assert wb["Transactions"].max_row == 1  # header only
    pdf = PdfReader(io.BytesIO((await _export(client, "pdf")).content))
    text = " ".join(p.extract_text() for p in pdf.pages)
    assert "Secret bonus" not in text and "No transactions match these filters." in text
